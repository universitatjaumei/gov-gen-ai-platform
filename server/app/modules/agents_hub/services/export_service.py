"""ExportService para workspaces de agentes — 1C.4.

Deploy: edge
Convierte final_document (Markdown) + run_manifest en DOCX/ODT descargable.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass
from typing import Any, Literal


@dataclass
class ExportResult:
    format: str
    file_path: str
    download_url: str
    size_bytes: int


class ExportService:
    """Genera DOCX/ODT a partir de un Markdown final y el run_manifest de la sesión."""

    def __init__(self, storage: Any) -> None:
        self._storage = storage

    async def export(
        self,
        workspace_id: str,
        final_document: str,
        run_manifest: Any,
        format: str = "docx",  # noqa: A002
        theme: Any = None,
    ) -> ExportResult:
        ext = format.lower()
        if ext == "docx":
            data = self._build_docx(final_document, run_manifest, theme)
        else:
            data = self._build_odt(final_document, run_manifest, theme)

        filename = f"informe_{workspace_id}.{ext}"
        file_path = f"exports/{workspace_id}/{filename}"
        await self._storage.put(file_path, data)

        return ExportResult(
            format=ext,
            file_path=file_path,
            download_url=file_path,
            size_bytes=len(data),
        )

    # ------------------------------------------------------------------
    # DOCX
    # ------------------------------------------------------------------

    def _build_docx(self, markdown: str, run_manifest: Any, theme: Any) -> bytes:
        from docx import Document

        doc = Document()
        font_family = _resolve_font(theme)

        for elem in _parse_markdown(markdown):
            if elem["type"] == "heading":
                level = min(elem["level"], 9)
                heading = doc.add_heading(elem["text"], level=level)
                if font_family:
                    for run in heading.runs:
                        run.font.name = font_family
            else:
                p = doc.add_paragraph(elem["text"])
                if font_family:
                    for run in p.runs:
                        run.font.name = font_family

        chunks = _get_chunks(run_manifest)
        if chunks:
            doc.add_heading("Referencias", level=1)
            for i, chunk in enumerate(chunks, 1):
                ref = f"[{i}] {getattr(chunk, 'source_url', '') or ''}"
                page = getattr(chunk, "page", None)
                if page is not None:
                    ref += f", p. {page}"
                doc.add_paragraph(ref)

        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    # ------------------------------------------------------------------
    # ODT
    # ------------------------------------------------------------------

    def _build_odt(self, markdown: str, run_manifest: Any, theme: Any) -> bytes:
        from odf.opendocument import OpenDocumentText
        from odf.text import H, P

        odt = OpenDocumentText()

        for elem in _parse_markdown(markdown):
            if elem["type"] == "heading":
                level = min(elem["level"], 10)
                odt.text.addElement(H(outlinelevel=level, text=elem["text"]))
            else:
                odt.text.addElement(P(text=elem["text"]))

        chunks = _get_chunks(run_manifest)
        if chunks:
            odt.text.addElement(H(outlinelevel=1, text="Referencias"))
            for i, chunk in enumerate(chunks, 1):
                ref = f"[{i}] {getattr(chunk, 'source_url', '') or ''}"
                page = getattr(chunk, "page", None)
                if page is not None:
                    ref += f", p. {page}"
                odt.text.addElement(P(text=ref))

        buf = io.BytesIO()
        odt.save(buf)
        return buf.getvalue()


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _parse_markdown(md: str) -> list[dict]:
    elements: list[dict] = []
    for line in md.splitlines():
        line = line.rstrip()
        if not line:
            continue
        m = re.match(r"^(#{1,6})\s+(.+)$", line)
        if m:
            elements.append({"type": "heading", "level": len(m.group(1)), "text": m.group(2)})
        else:
            text = re.sub(r"\[\^\d+\]", "", line).strip()
            if text:
                elements.append({"type": "paragraph", "text": text})
    return elements


def _get_chunks(run_manifest: Any) -> list:
    return list(getattr(run_manifest, "retrieved_chunks", None) or [])


def _resolve_font(theme: Any) -> str | None:
    if theme is None:
        return None
    try:
        return theme.typography.get("font_family")
    except (AttributeError, TypeError):
        return None
