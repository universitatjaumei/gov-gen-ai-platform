"""Fixtures para los tests del módulo redacción (9R).

Lleva además el mismo orden de importación que sus hermanos de `agents_hub` y `curation`, por la
razón que está escrita abajo y que no es opcional en Windows.
"""
from __future__ import annotations

# ── Orden de importación obligatorio, no reordenar ─────────────────────────────
# `langchain_text_splitters` arrastra `sentence_transformers` → torch. En Windows, cargar torch
# DESPUÉS de haber abierto una conexión asyncpg aborta el proceso con «Windows fatal exception:
# access violation»; al revés funciona. Diagnosticado el 2026-07-29 y reducido a dos líneas:
#
#     asyncpg connect  →  import langchain_text_splitters   ⇒ access violation
#
# Este fichero lo perdió al ganar su `db_session` —los dos hermanos sí lo tienen—, así que la
# suite de redacción podía tumbar el proceso en la plataforma de desarrollo de este proyecto.
# Lo señaló la revisión automática de la PR del despliegue (APER.16).
import langchain_text_splitters  # noqa: F401  ← debe ir primero

from pathlib import Path

import pytest


@pytest.fixture(scope="module")
def pdf_text_pipeline():
    """Instancia única de PDFTextExtractionPipeline por módulo.

    Scope=module evita recargar los modelos Docling en cada test.
    """
    from server.app.modules.redaccion.pipelines.pdf_text_pipeline import PDFTextExtractionPipeline
    return PDFTextExtractionPipeline()


@pytest.fixture
def sample_xlsx(tmp_path: Path) -> Path:
    """Excel con una sola hoja: columnas nombre/valor/categoria, 3 filas de datos."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["nombre", "valor", "categoria"])
    ws.append(["Alpha", 100, "A"])
    ws.append(["Beta", 200, "B"])
    ws.append(["Gamma", 300, "A"])
    path = tmp_path / "sample.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def multi_sheet_xlsx(tmp_path: Path) -> Path:
    """Excel con dos hojas de esquema diferente."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Sheet1"
    ws1.append(["nombre", "valor"])
    ws1.append(["Alpha", 100])

    ws2 = wb.create_sheet("Sheet2")
    ws2.append(["ciudad", "poblacion"])
    ws2.append(["Valencia", 800000])
    ws2.append(["Madrid", 3400000])

    path = tmp_path / "multi_sheet.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def sample_pdf(tmp_path: Path) -> Path:
    """PDF digital con capa de texto extractable."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas as rl_canvas

    path = tmp_path / "sample.pdf"
    c = rl_canvas.Canvas(str(path), pagesize=A4)
    c.drawString(100, 750, "Primer parrafo del informe de analisis.")
    c.drawString(100, 720, "Datos del periodo: enero-marzo 2026.")
    c.drawString(100, 690, "Conclusion: resultados satisfactorios.")
    c.save()
    return path


@pytest.fixture
def image_only_pdf(tmp_path: Path) -> Path:
    """PDF sin capa de texto (solo graficos vectoriales), no extractable."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas as rl_canvas

    path = tmp_path / "image_only.pdf"
    c = rl_canvas.Canvas(str(path), pagesize=A4)
    c.setFillColorRGB(0.9, 0.9, 0.9)
    c.rect(50, 50, 495, 740, fill=1, stroke=0)
    c.save()
    return path


@pytest.fixture
def pdf_with_tables(tmp_path: Path) -> Path:
    """PDF con tabla estructurada de 3 columnas y 4 filas, detectable por camelot."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

    path = tmp_path / "with_tables.pdf"
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=50, rightMargin=50)
    data = [
        ["Concepto", "Importe", "Porcentaje"],
        ["Ingresos", "10000 EUR", "100%"],
        ["Gastos", "7500 EUR", "75%"],
        ["Beneficio", "2500 EUR", "25%"],
    ]
    tbl = Table(data, colWidths=[250, 160, 115])
    tbl.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
    ]))
    doc.build([tbl])
    return path


@pytest.fixture
async def db_session(db_url: str):
    """Sesión sobre la BD desechable, para los tests de redacción que tocan tablas.

    Copia deliberada de la de `tests/modules/curation/conftest.py`, por el mismo motivo que
    aquélla lo era de la de `agents_hub`: las entidades viven en `HubOperationalBase` y necesitan
    el mismo motor. Sin `drop_all`: la BD entera se borra al final del test.
    """
    from server.app.modules.agents_hub.database.connection import (
        create_async_engine,
        create_session_factory,
    )

    engine = create_async_engine(db_url)
    session_factory = create_session_factory(engine)
    async with session_factory() as session:
        yield session
    await engine.dispose()
