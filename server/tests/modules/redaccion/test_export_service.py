"""Tests 9R.9.2 — ExportService DOCX (RED → GREEN).

Alcance MVP: DOCX únicamente. ODT en backlog post-MVP.
Deploy: edge
"""
from __future__ import annotations

import io
import shutil
import uuid
from datetime import datetime, timezone

import pytest

from server.app.modules.redaccion.contracts.manifest import (
    AIBlockSummary,
    DraftingRunManifest,
    UploadedDocumentInfo,
)
from server.app.modules.redaccion.contracts.runtime import (
    ApprovalRecord,
    ExtractionWarning,
)
from server.app.modules.redaccion.services.export_service import (
    ExportNotReadyError,
    ExportService,
)

# ---------------------------------------------------------------------------
# Factories
# ---------------------------------------------------------------------------

_HASH = "sha256:abc123deadbeef"
_CONTENT = "Primer párrafo del informe.\nSegundo párrafo con datos relevantes."


def _make_manifest(**kwargs) -> DraftingRunManifest:
    defaults = dict(
        id=uuid.uuid4(),
        workspace_id=uuid.uuid4(),
        template_id=uuid.uuid4(),
        template_version_id=uuid.uuid4(),
        report_profile="GENERIC_REPORT",
        status_at_close="assembled",
        final_document_hash=_HASH,
    )
    defaults.update(kwargs)
    return DraftingRunManifest(**defaults)


def _open_docx(raw: bytes):
    from docx import Document
    return Document(io.BytesIO(raw))


def _all_text(doc) -> str:
    parts = [p.text for p in doc.paragraphs]
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


# ---------------------------------------------------------------------------
# Tests
# ---------------------------------------------------------------------------

def test_export_service_reads_run_manifest_before_export():
    manifest = _make_manifest()
    svc = ExportService()

    raw = svc.export_to_docx(manifest, _CONTENT)

    doc = _open_docx(raw)
    text = _all_text(doc)
    # Hash from manifest is present in the DOCX
    assert _HASH in text
    # Content is included
    assert "Primer párrafo" in text


def test_export_service_rejects_export_without_final_hash():
    manifest = _make_manifest(final_document_hash=None)
    svc = ExportService()

    with pytest.raises(ExportNotReadyError):
        svc.export_to_docx(manifest, _CONTENT)


def test_export_appendix_includes_ai_blocks_with_model():
    ai_blocks = [
        AIBlockSummary(
            block_id="b_intro",
            kind="AI_ASSISTED_TEXT",
            status="approved",
            model_used="claude-opus-4-7",
            prompt_version="v2",
        ),
        AIBlockSummary(
            block_id="b_summary",
            kind="AI_SUMMARY",
            status="approved",
            model_used="claude-sonnet-4-6",
            prompt_version="v1",
        ),
    ]
    manifest = _make_manifest(ai_blocks=ai_blocks)
    svc = ExportService()

    raw = svc.export_to_docx(manifest, _CONTENT)
    doc = _open_docx(raw)
    text = _all_text(doc)

    assert "b_intro" in text
    assert "claude-opus-4-7" in text
    assert "v2" in text
    assert "b_summary" in text
    assert "claude-sonnet-4-6" in text


def test_export_appendix_includes_user_approvals():
    approver_id = uuid.uuid4()
    approvals = [
        ApprovalRecord(
            approved_by=approver_id,
            approved_at=datetime(2026, 5, 17, 10, 0, 0, tzinfo=timezone.utc),
            note="Conforme con el análisis",
        )
    ]
    manifest = _make_manifest(user_approvals=approvals)
    svc = ExportService()

    raw = svc.export_to_docx(manifest, _CONTENT)
    doc = _open_docx(raw)
    text = _all_text(doc)

    assert str(approver_id) in text
    assert "Conforme con el análisis" in text
    assert "2026-05-17" in text


def test_final_document_hash_in_export_metadata_matches_manifest():
    manifest = _make_manifest(final_document_hash=_HASH)
    svc = ExportService()

    raw = svc.export_to_docx(manifest, _CONTENT)
    doc = _open_docx(raw)
    text = _all_text(doc)

    assert _HASH in text


def test_docx_opens_without_errors_with_python_docx():
    manifest = _make_manifest(
        uploaded_documents=[
            UploadedDocumentInfo(
                slot_id="slot_1",
                filename="datos.xlsx",
                storage_path="s/datos.xlsx",
                size_bytes=20480,
                uploaded_at=datetime.now(timezone.utc),
            )
        ],
        ai_blocks=[
            AIBlockSummary(
                block_id="b_intro",
                kind="AI_ASSISTED_TEXT",
                status="approved",
                model_used="claude-opus-4-7",
            )
        ],
        warnings=[
            ExtractionWarning(
                block_id="b_data",
                message="Confianza baja en columna 'importe'",
                kind="low_confidence",
            )
        ],
    )
    svc = ExportService()

    raw = svc.export_to_docx(manifest, _CONTENT)

    # Must open cleanly and contain expected sections
    doc = _open_docx(raw)
    text = _all_text(doc)

    assert "Anexo de Auditoría" in text
    assert "datos.xlsx" in text
    assert "b_intro" in text
    assert "Confianza baja" in text
    assert len(raw) > 0


@pytest.mark.skipif(
    shutil.which("libreoffice") is None and shutil.which("soffice") is None,
    reason="LibreOffice not available in PATH",
)
def test_docx_renders_in_libreoffice_when_available(tmp_path):
    """Converts the DOCX to PDF via LibreOffice to verify render compatibility."""
    import subprocess

    manifest = _make_manifest()
    svc = ExportService()
    raw = svc.export_to_docx(manifest, _CONTENT)

    docx_path = tmp_path / "informe.docx"
    docx_path.write_bytes(raw)

    lo = shutil.which("libreoffice") or shutil.which("soffice")
    result = subprocess.run(
        [lo, "--headless", "--convert-to", "pdf", "--outdir", str(tmp_path), str(docx_path)],
        capture_output=True,
        timeout=60,
    )
    assert result.returncode == 0
    pdf_files = list(tmp_path.glob("*.pdf"))
    assert len(pdf_files) == 1
    assert pdf_files[0].stat().st_size > 0
