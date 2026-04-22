# client_app/app/core/exporters.py
"""
Data Exporters for Client (On-Premise).

Provides Excel export functionality with data humanization and traceability.
"""
import os
import pandas as pd
from datetime import datetime
from typing import List, Dict, Any, Tuple, Optional

from openpyxl import Workbook
from openpyxl.comments import Comment


def _humanizar_valor(valor: Any) -> str:
    """
    Convierte cualquier estructura (dict, list, primitivo) a una cadena limpia para Excel.
    Logica:
      1. Limpieza: Elimina claves tecnicas de diccionarios.
      2. Flattening: Si queda 1 clave, usa su valor. Si hay varias, 'K: V | K2: V2'.
      3. Listas: Procesa cada elemento y unelos con ", ".
      4. Primitivos: str(valor).
      5. None/Vacio: "".
    """
    if valor is None:
        return ""

    # Caso recursivo para listas
    if isinstance(valor, list):
        # Procesar cada elemento y filtrar vacios
        items = [_humanizar_valor(x) for x in valor]
        return ", ".join(filter(None, items))

    # Caso Diccionarios
    if isinstance(valor, dict):
        # 1. Copia y Limpieza
        d = valor.copy()
        keys_to_remove = ['pagina', 'page', 'coincidencia_exacta', 'confidence', 'valor_original', 'source']
        for k in keys_to_remove:
            d.pop(k, None)

        # Si tras limpiar esta vacio, retornar cadena vacia
        if not d:
            return ""

        # 2. Smart Flattening
        if len(d) == 1:
            # Devuelve solo el valor de la unica clave
            return str(list(d.values())[0])
        else:
            # Formato compacto 'Key: Val | Key2: Val2'
            # Capitalizamos claves para mejor lectura 'nombre' -> 'Nombre'
            parts = [f"{str(k).capitalize()}: {v}" for k, v in d.items()]
            return " | ".join(parts)

    # Primitivos
    return str(valor)


def formatear_excel_dual(resultados_raw: List[Dict[str, Any]], nombre_archivo_base: str = "extraccion_LLM", output_dir: Optional[str] = None) -> Tuple[str, pd.DataFrame]:
    """
    Genera un Excel consolidado (una sola hoja).
    - Valor: Limpio y humanizado.
    - Comentario: Trazabilidad (Origen, Pagina, Confianza).
    """
    if not output_dir:
        output_dir = "data/resultados"

    if not os.path.exists(output_dir):
        os.makedirs(output_dir)

    wb = Workbook()
    ws = wb.active
    ws.title = "Extraccion Consolidada"

    # 1. Recopilar todas las claves (columnas) posibles
    #    Ignorando tecnicas
    TECH_KEYS_TO_EXCLUDE = {'status', 'processing_seconds', 'timestamp', 'cost', 'token_usage', 'error', 'meta', 'processing_metadata', 'filename', 'archivo_origen'}

    def _es_columna_valida(key: str) -> bool:
        k = key.lower()
        if k in TECH_KEYS_TO_EXCLUDE:
            return False
        if key.startswith(('IGNORED_', 'META_', '_')):
            return False
        return True

    all_keys = set()
    for item in resultados_raw:
        data = item.get('data', {}) or item.get('datos', {})

        # [PATCH] Handle List of items (Multi-row per file)
        if isinstance(data, list):
            for sub_item in data:
                if isinstance(sub_item, dict):
                    for k in sub_item.keys():
                        if _es_columna_valida(k):
                            all_keys.add(k)
        elif isinstance(data, dict):
            for k in data.keys():
                if _es_columna_valida(k):
                    all_keys.add(k)

    # Ordenar columnas: Archivo primero, luego alfabetico
    headers = ['Archivo'] + sorted(list(all_keys))
    ws.append(headers)

    # Preparamos DataFrame para retorno (solo valores limpios)
    rows_for_df = []

    for item in resultados_raw:
        archivo = item.get('filename') or item.get('archivo_origen', "Desconocido")
        status = item.get('status', 'ok')
        raw_data = item.get('data', {}) or item.get('datos', {})

        # Helper to process a single data dict row
        def process_row_data(d_row):
            row_dict = {'Archivo': archivo}
            excel_row = [archivo]
            comments_map = {}

            for i, header in enumerate(headers[1:], start=1):
                val_obj = d_row.get(header)
                clean_val = ""
                comment_str = ""

                if val_obj is not None:
                    if isinstance(val_obj, dict) and 'valor' in val_obj:
                        raw_val = val_obj.get('valor')

                        # Normalize None -> "" (User Request)
                        if raw_val is None:
                            clean_val = ""
                            # If explicitly None in dict, it likely means checked but empty/optional
                            if not comment_str:
                                comment_str = "Opcional no presente"
                        else:
                            clean_val = _humanizar_valor(raw_val)

                        page = val_obj.get('pagina')
                        source = val_obj.get('source', '')
                        conf = val_obj.get('confidence')

                        details = []
                        if source:
                            details.append(f"Fuente: {'LLM (snippet)' if source=='llm' else 'Script'}")
                        if page is not None:
                            details.append(f"Pag: {page}")
                        if conf is not None:
                            details.append(f"Conf: {conf:.2f}")

                        # Add specific note if we marked it
                        if comment_str == "Opcional no presente":
                            details.insert(0, "[Vacio/Opcional]")

                        if details:
                            comment_str = "\n".join(details)
                    else:
                        clean_val = _humanizar_valor(val_obj)
                else:
                    # Header exists globally but missing in this row
                    clean_val = ""
                    comment_str = "Opcional no presente"

                excel_row.append(clean_val)
                row_dict[header] = clean_val
                if comment_str:
                    comments_map[i] = comment_str

            ws.append(excel_row)
            current_row_idx = ws.max_row
            for col_idx, note in comments_map.items():
                cell = ws.cell(row=current_row_idx, column=col_idx + 1)
                cell.comment = Comment(note, f"Gov Gen AI ({archivo})")

            rows_for_df.append(row_dict)

        # [PATCH] Explode List Data
        if isinstance(raw_data, list):
            for sub_data in raw_data:
                if isinstance(sub_data, dict):
                    process_row_data(sub_data)
        elif isinstance(raw_data, dict):
            process_row_data(raw_data)

    # Save
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    nombre_excel = os.path.join(output_dir, f"{nombre_archivo_base}_{timestamp}.xlsx")
    wb.save(nombre_excel)

    return nombre_excel, pd.DataFrame(rows_for_df)
